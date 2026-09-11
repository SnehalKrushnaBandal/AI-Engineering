import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
import time

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API Error")

client = Groq(api_key=my_api_key)
model = "openai/gpt-oss-120b"

# ********* Part 1: Job Parsing
# 1
job_description = """
We are looking for a Junior Python Developer to join our software development team.

The ideal candidate should have a Bachelor's degree in Computer Science, Information Technology, or a related field and 1-2 years of experience in software development. Freshers with strong programming skills and relevant projects may also apply.

Required Skills:
• Strong knowledge of Python programming
• Good understanding of Object-Oriented Programming
• Experience with REST APIs
• Knowledge of SQL and relational databases
• Familiarity with Git and GitHub
• Basic understanding of data structures and algorithms
• Good problem-solving and debugging skills

Preferred Skills:
• Experience with Flask or FastAPI
• Knowledge of PostgreSQL or MySQL
• Familiarity with Docker
• Basic knowledge of cloud platforms such as AWS or Google Cloud
• Experience working with React.js is a plus

Key Responsibilities:
• Develop and maintain Python-based applications and backend services.
• Design and develop RESTful APIs for web applications.
• Work with SQL databases to store, retrieve, and manage application data.
• Write clean, maintainable, and well-documented code.
• Debug and resolve software issues.
• Participate in code reviews and technical discussions.
• Collaborate with frontend developers and other team members.
• Use Git and GitHub for source code management.
• Write unit tests and perform application testing.
• Learn and adopt new technologies as required by the project.

Education:
Bachelor's degree in Computer Science, Information Technology, Software Engineering, or a related field.

Experience:
1-2 years of professional software development experience. Freshers with strong academic projects may also be considered.
"""

# 2
class Job_Desc_Class(BaseModel):
    role : str
    required_skills : list[str]
    preferred_skills : list[str]
    minimum_experience : float | None
    education_requirements : list[str]
    responsibilities : list[str]

# 3
job_desc_schema = Job_Desc_Class.model_json_schema()

# 4 
system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job description and extract structured information from them.

Return only valid JSON matching this schema : {job_desc_schema}

IMPORTANT: 
DO NOT return the schema itself.
DO NOT return fields like "properties", "type" or "title".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.

If information for a list is missing, return an empty list.
Do not invent information.
"""

# 5
user_prompt = f"""
Analyze the following job description:
{job_description}
"""

# 6 
system_msg = {
    "role" : "system",
    "content" : system_prompt
}

# 7
user_msg = {
    "role" : "user",
    "content" : user_prompt
}

# 8
response_format = {
    "type" : "json_object"
}

# 9
messages = [system_msg, user_msg]

# 10
response = client.chat.completions.create(model = model, messages = messages, response_format = response_format)

answer = response.choices[0].message.content

# 11 Parse the JSON response into the Pydantic model
import json
raw_json = answer
job_data = json.loads(raw_json)
job = Job_Desc_Class(**job_data)

# print(job.minimum_experience)
# print(job.education_requirements)

# ********* Part 2 : Resume Parsing 

# 1 : This experience is created just to use in Resume_Class
class Experience_Class(BaseModel):
    company : str | None = None
    role : str | None = None
    duration : str | None = None
    description : str | None = None
    skills_used : list[str] = []

# 2 : Main Resume_Class
class Resume_Class(BaseModel):
    name : str | None = None
    email : str | None = None
    phone : str | None = None

    total_experience_years : float | None = None
    skills : list[str] = []
    experiences : list[Experience_Class] = []
    education : list[str] = []
    projects : list[str] = []
    certifications : list[str] = []

# 3
resume_schema = Resume_Class.model_json_schema()


# ******** Part 4 functions
# 7
def parse_resume(resume_text):

    system_prompt = f"""

You are an expert resume parser.

Extract information from the resume based on its meaning, not only based on exact section headings.

Different resumes may use different headings.

For example: 
- Experince
- Professional Experince
- Work History
- Employment
- Internships

These may all contain relevant experince.

Skills may appear in:
- Skills section
- Work experience
- Internships
- Projects

Return ONLY valid JSON matching this schema: 
{resume_schema} 

Important rules: 

1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
6. Return ONLY the JSON object.
7. Do not use markdown.
8. Do not add explanations before or after the JSON.
"""

    # 8
    user_prompt = f"""
        Parse the following resume: 
        {resume_text}
    """

    # 9
    system_msg = {
        "role" : "system",
        "content" : system_prompt
    }

    # 10
    user_msg = {
        "role" : "user",
        "content" : user_prompt
    }

    # 11
    messages = [system_msg, user_msg]

    response_format = {
        "type" : "json_object"
    }

    for attempt in range(3):

        try:

            response = client.chat.completions.create(
                model=model,
                messages=messages,
                response_format=response_format,
                temperature=0
            )

            raw_output = response.choices[0].message.content

            data = json.loads(raw_output)

            resume = Resume_Class(**data)

            return resume

        except Exception as e:

            print(f"Attempt {attempt + 1} failed: {e}")

            if attempt == 2:
                raise

            time.sleep(2)

    # # 12
    # response = client.chat.completions.create(model = model, messages = messages, response_format = response_format)

    # # 13
    # raw_output = response.choices[0].message.content
    # # print("\nRAW RESUME OUTPUT:")
    # # print(raw_output)

    # data = json.loads(raw_output)
    # resume = Resume_Class(**data)
    # return resume



# 16 final_score method

class MatchDetails_Class(BaseModel):
    candidate_name: str
    matching_skills: list[str]
    missing_skills: list[str]
    experience_requirement_met: bool
    overall_match_percentage: float
    final_verdict: str

class MatchResult_Class(BaseModel):
    score : float
    details : MatchDetails_Class


def final_score(job, resume):


    prompt = f"""
You are an HR recruiter.

Compare the candidate resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json()}

CANDIDATE RESUME:
{resume.model_dump_json()}

Return ONLY valid JSON with exactly this structure:

{{
    "score": 0,
    "details": {{
        "candidate_name": "",
        "matching_skills": [],
        "missing_skills": [],
        "experience_requirement_met": false,
        "overall_match_percentage": 0,
        "final_verdict": ""
    }}
}}

Rules:

1. score must be between 0 and 100.
2. overall_match_percentage must be between 0 and 100.
3. matching_skills must contain only skills actually present in the resume.
4. missing_skills must contain important skills required by the job but missing from the resume.
5. Do not invent candidate experience or skills.
6. Keep final_verdict short.
7. Return ONLY JSON.
"""

    messages = [
        {
            "role": "user",
            "content": prompt
        }
    ]

    response_format = {
        "type": "json_object"
    }

    for attempt in range(3):

        try:

            response = client.chat.completions.create(
                model=model,
                messages=messages,
                response_format=response_format,
                temperature=0
            )

            raw_output = response.choices[0].message.content

            data = json.loads(raw_output)

            return MatchResult_Class(**data)

        except Exception as e:

            print(f"Scoring attempt {attempt + 1} failed: {e}")

            if attempt == 2:
                raise

            wait_time = 2 ** attempt
            print(f"Retrying in {wait_time} seconds...")

            time.sleep(wait_time)

# ********* Part 3 : Reading resume
from pypdf import PdfReader
from docx import Document

# 1 function to read pdf file resume
def read_pdf_resume(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if(page_text):
            text += page_text + "\n"
    return text

# 2 function to read docx file resume
def read_docx_resume(file_path):
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text

def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf_resume(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx_resume(file_path)
    else:
        return None


# ******** Part 4 : Actual Resume Parsing
    
# 1
resume_folder = Path("resumes")

# 2
all_result = []

# 3
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
# 4
    print("\nProcessing: ", file_path.name)
    # 5
    resume_text = read_resume(file_path)

    # 6
    parsed_resume = parse_resume(resume_text)  # now check parse_resume method

    # 14
    time.sleep(5)

    # 15 
    result = final_score(job, parsed_resume)
    time.sleep(5)

    # 17
    print("Score: ", result.score)
    all_result.append({
        "name" : parsed_resume.name,
        "score" : result.score,
        "details" : result.details
    })

# 18
all_result.sort(
    key = lambda candidate: candidate["score"],
    reverse = True
)

# 19
# top_2 = all_result[:2]
# worst_2 = all_result[-2:]

top_2 = all_result[:2]

if len(all_result) >= 5:
    worst_2 = all_result[-2:]
else:
    worst_2 = []

# 20
print("\nTOP 2 CANDIDATES")
for candidate in top_2:
    # print(
    #     candidate["name"],
    #     "-",
    #     candidate["score"],
    #     "%"     
    # )
    details = candidate["details"]
    print("\n----------------------------------------")
    print("Name:", details.candidate_name)
    print("Score:", details.overall_match_percentage, "%")
    print("Matching Skills:", details.matching_skills)
    print("Missing Skills:", details.missing_skills)
    print("Experience Met:", details.experience_requirement_met)
    print("Verdict:", details.final_verdict)
    print("----------------------------------------")

    # print(candidate["details"], "\n")

# 21
print("\nLOWEST 2 CANDIDATES")
if worst_2:

    for candidate in worst_2:
        # print(
        #     candidate["name"],
        #     "-",
        #     candidate["score"],
        #     "%"     
        # )
        details = candidate["details"]
        print("\n----------------------------------------")
        print("Name:", details.candidate_name)
        print("Score:", details.overall_match_percentage, "%")
        print("Matching Skills:", details.matching_skills)
        print("Missing Skills:", details.missing_skills)
        print("Experience Met:", details.experience_requirement_met)
        print("Verdict:", details.final_verdict)
        print("----------------------------------------\n")
        # print(candidate["details"], "\n")
else:
    print("Not enough candidates to display separately.")